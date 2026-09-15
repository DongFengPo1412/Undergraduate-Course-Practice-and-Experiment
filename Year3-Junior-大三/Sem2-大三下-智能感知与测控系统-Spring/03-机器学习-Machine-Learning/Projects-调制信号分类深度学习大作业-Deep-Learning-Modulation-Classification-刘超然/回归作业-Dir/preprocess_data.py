import os
import openpyxl
import numpy as np

def load_battery_data(file_path):
    print(f"Loading battery data from {os.path.basename(file_path)}...")
    wb = openpyxl.load_workbook(file_path, read_only=True)
    sheet = wb.active
    
    # Read headers
    iterator = sheet.iter_rows(max_row=1)
    headers = [cell.value for cell in next(iterator)]
    
    # Map headers to indices
    cc_idx = headers.index('cc_duration')
    cv_idx = headers.index('cv_duration')
    time_idx = headers.index('timeInterval')
    ica_idx = headers.index('ica_max')
    soh_idx = headers.index('SOH')
    
    X = []
    y = []
    for row in sheet.iter_rows(min_row=2):
        vals = [cell.value for cell in row]
        # Skip if row is empty or incomplete
        if len(vals) <= max(cc_idx, cv_idx, time_idx, ica_idx, soh_idx):
            continue
        if any(v is None for v in [vals[cc_idx], vals[cv_idx], vals[time_idx], vals[ica_idx], vals[soh_idx]]):
            continue
        X.append([vals[cc_idx], vals[cv_idx], vals[time_idx], vals[ica_idx]])
        y.append(vals[soh_idx])
        
    wb.close()
    return np.array(X, dtype=np.float32), np.array(y, dtype=np.float32)

def main():
    data_dir = r"C:\Users\Liu\Desktop\大三下学期文件\机器学习\回归作业\CLACEdata"
    scratch_dir = os.path.dirname(os.path.abspath(__file__))
    
    train_files = [
        "CS2_33.xlsx", "CS2_34.xlsx", "CS2_35.xlsx", 
        "CS2_36.xlsx", "CS2_37.xlsx", "CS2_38.xlsx",
        "CX2_34.xlsx"
    ]
    test_file = "CX2_36.xlsx"
    
    # 1. Load Train Data
    X_train_list = []
    y_train_list = []
    
    for f in train_files:
        fp = os.path.join(data_dir, f)
        if os.path.exists(fp):
            X_f, y_f = load_battery_data(fp)
            X_train_list.append(X_f)
            y_train_list.append(y_f)
        else:
            print(f"Warning: Train file not found: {fp}")
            
    X_train = np.vstack(X_train_list)
    y_train = np.concatenate(y_train_list)
    
    # 2. Load Test Data
    test_fp = os.path.join(data_dir, test_file)
    if os.path.exists(test_fp):
        X_test, y_test = load_battery_data(test_fp)
    else:
        print(f"Error: Test file not found: {test_fp}")
        return
        
    print(f"Loaded Train set size: {X_train.shape[0]} samples")
    print(f"Loaded Test set size: {X_test.shape[0]} samples")
    
    # 3. Export to NPZ file
    npz_path = os.path.join(scratch_dir, "battery_data.npz")
    np.savez(npz_path, X_train=X_train, y_train=y_train, X_test=X_test, y_test=y_test)
    print(f"Data successfully saved to {npz_path}")

if __name__ == "__main__":
    main()
